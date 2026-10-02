#!/usr/bin/env python3
"""Generate branded VC Retail Analytics PPTX — visuals, deep wiring (local+Snowflake), clickable demo URL."""

from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

BG = RGBColor(0xF3, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF8)
INK = RGBColor(0x1C, 0x19, 0x14)
MUTED = RGBColor(0x6B, 0x64, 0x58)
ACCENT = RGBColor(0x8A, 0x6A, 0x3B)
LINE = RGBColor(0xD9, 0xD0, 0xC0)
DARK = RGBColor(0x1C, 0x19, 0x14)
LIGHT = RGBColor(0xF4, 0xEF, 0xE4)
SOFT = RGBColor(0xC8, 0xC0, 0xB0)
WARM = RGBColor(0xE8, 0xDF, 0xCF)

ASSETS = Path(__file__).resolve().parents[1] / "docs" / "deck" / "assets"
DEMO_URL = os.environ.get("VC_DEMO_URL", "https://vc.datasharkbi.com/").rstrip("/") + "/"
TOTAL = 17


def set_run(run, size=18, bold=False, color=INK, font="Georgia", underline=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.underline = underline


def add_bg(slide, w, h, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp = shape._element
    slide.shapes._spTree.remove(sp)
    slide.shapes._spTree.insert(2, sp)


def add_rect(slide, left, top, width, height, fill=CARD, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, font="Georgia", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_link_text(slide, left, top, width, height, label, url, size=16, color=ACCENT, align=PP_ALIGN.LEFT, bold=True):
    """Text run hyperlink (PowerPoint)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = label
    set_run(run, size=size, bold=bold, color=color, font="Calibri", underline=True)
    run.hyperlink.address = url
    return box


def add_link_button(slide, left, top, width, height, label, url, fill=ACCENT, text_color=LIGHT, size=16):
    """Shape click-action hyperlink — works in PowerPoint / Keynote more reliably than text alone."""
    shape = add_rect(slide, left, top, width, height, fill=fill)
    shape.click_action.hyperlink.address = url
    # Centered label on top of button (also linked)
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        # vertical center via bodyPr
        body_pr = tf._txBody.bodyPr
        body_pr.set("anchor", "ctr")
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    set_run(run, size=size, bold=True, color=text_color, font="Calibri", underline=False)
    # Also attach hlink on the run as belt-and-suspenders
    run.hyperlink.address = url
    return shape


def add_bullets(slide, left, top, width, height, items, color=MUTED, size=13):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0:
            p.clear()
        p.space_before = Pt(5)
        run = p.add_run()
        run.text = "▪  " + item
        set_run(run, size=size, color=color, font="Calibri")
    return box


def footer(slide, page):
    add_textbox(slide, Inches(0.5), Inches(7.12), Inches(6.5), Inches(0.28), "VC Retail Analytics  ·  synthetic", size=10, color=MUTED, font="Calibri")
    # Small clickable pill in footer
    add_link_button(slide, Inches(7.2), Inches(7.05), Inches(4.3), Inches(0.32), DEMO_URL.rstrip("/"), DEMO_URL, fill=WARM, text_color=ACCENT, size=10)
    add_textbox(slide, Inches(11.7), Inches(7.12), Inches(1.2), Inches(0.28), f"{page}/{TOTAL}", size=10, color=MUTED, font="Calibri", align=PP_ALIGN.RIGHT)


def eyebrow(slide, text):
    add_textbox(slide, Inches(0.65), Inches(0.32), Inches(11), Inches(0.28), text.upper(), size=11, bold=True, color=ACCENT, font="Calibri")


def title_bar(slide, text):
    add_rect(slide, Inches(0.65), Inches(0.72), Inches(0.08), Inches(0.48), fill=ACCENT)
    add_textbox(slide, Inches(0.9), Inches(0.68), Inches(11.5), Inches(0.55), text, size=28, color=INK, font="Georgia")


def pic(slide, name, left, top, width, height=None):
    path = ASSETS / name
    if not path.exists():
        return None
    if height is None:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def build(out_path: str) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height

    # ----- 1 Title -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H, DARK)
    pic(s, "vc-hero-lighting.jpg", Inches(5.5), 0, Inches(7.85), Inches(7.5))
    add_rect(s, Inches(5.5), 0, Inches(0.08), H, fill=ACCENT)
    add_textbox(s, Inches(0.6), Inches(1.15), Inches(4.6), Inches(0.35), "VISUAL COMFORT–FLAVORED  ·  PUBLIC DEMO", size=11, bold=True, color=ACCENT, font="Calibri")
    add_textbox(s, Inches(0.6), Inches(1.65), Inches(4.7), Inches(1.8), "Retail analytics\nwiring,\nexplained.", size=36, color=LIGHT, font="Georgia")
    add_textbox(s, Inches(0.6), Inches(3.7), Inches(4.6), Inches(0.9), "Local SQLite + optional Snowflake — same UI, governed metrics, synthetic dealers.", size=14, color=SOFT, font="Calibri")
    add_link_button(s, Inches(0.6), Inches(4.85), Inches(4.4), Inches(0.7), "Open live demo →", DEMO_URL, fill=ACCENT, size=18)
    add_link_text(s, Inches(0.6), Inches(5.7), Inches(4.6), Inches(0.35), DEMO_URL, DEMO_URL, size=13, color=SOFT, bold=False)
    add_textbox(s, Inches(0.6), Inches(6.15), Inches(4.6), Inches(0.4), "Login: dlr-0001 / vc-demo", size=13, color=SOFT, font="Calibri")
    add_textbox(s, Inches(0.6), Inches(6.6), Inches(4.6), Inches(0.3), "Paul Rendleman", size=13, color=LIGHT, font="Georgia")

    # ----- 2 Clickable demo -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Start here")
    title_bar(s, "Click the button — opens the live demo")
    add_link_button(s, Inches(0.65), Inches(1.55), Inches(12.0), Inches(1.15), DEMO_URL, DEMO_URL, fill=DARK, text_color=LIGHT, size=28)
    add_textbox(s, Inches(0.65), Inches(2.9), Inches(12), Inches(0.35), "If the button doesn’t open: copy the URL, or Ctrl/Cmd-click the underlined link below.", size=12, color=MUTED, font="Calibri")
    add_link_text(s, Inches(0.65), Inches(3.25), Inches(12), Inches(0.35), DEMO_URL, DEMO_URL, size=16, color=ACCENT)
    for i, (n, t, b) in enumerate([
        ("1", "Open", "Click the dark bar above"),
        ("2", "Sign in", "dlr-0001 / vc-demo"),
        ("3", "Confirm", "SYNTHETIC badge on"),
        ("4", "Walk", "Overview → Analytics → Assistant → Catalog"),
    ]):
        left = Inches(0.65) + Inches(i * 3.15)
        add_rect(s, left, Inches(3.9), Inches(3.0), Inches(2.7), fill=CARD, line=LINE)
        add_rect(s, left, Inches(3.9), Inches(3.0), Inches(0.08), fill=ACCENT)
        add_textbox(s, left + Inches(0.2), Inches(4.2), Inches(2.6), Inches(0.4), n, size=22, bold=True, color=ACCENT, font="Calibri")
        add_textbox(s, left + Inches(0.2), Inches(4.7), Inches(2.6), Inches(0.4), t, size=18, color=INK, font="Georgia")
        add_textbox(s, left + Inches(0.2), Inches(5.25), Inches(2.6), Inches(1.0), b, size=13, color=MUTED, font="Calibri")
    footer(s, 2)

    # ----- 3 Opportunity -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Why this exists")
    title_bar(s, "Support pattern vs warehouse truth")
    add_textbox(s, Inches(0.65), Inches(1.45), Inches(7.0), Inches(1.0), "Site support handles FAQ and escalation. Dealer sell-through, margin, and stock risk are not in chat — this demo grounds those answers in governed metrics (local or Snowflake).", size=15, color=MUTED, font="Calibri")
    pic(s, "vc-product-still.jpg", Inches(8.0), Inches(1.4), Inches(4.7), Inches(3.3))
    for i, (t, b) in enumerate([
        ("Session tenancy", "Dealer scope from auth cookie — never from the prompt."),
        ("Dual backend", "Same UI: SQLite local or Snowflake SERVING views."),
        ("Traceable assistant", "Proposed answers include SQL + params in the tool trace."),
    ]):
        left = Inches(0.65) + Inches(i * 4.05)
        add_rect(s, left, Inches(4.95), Inches(3.85), Inches(1.7), fill=CARD, line=LINE)
        add_textbox(s, left + Inches(0.2), Inches(5.15), Inches(3.45), Inches(0.4), t, size=15, color=INK, font="Georgia")
        add_textbox(s, left + Inches(0.2), Inches(5.6), Inches(3.45), Inches(0.85), b, size=12, color=MUTED, font="Calibri")
    footer(s, 3)

    # ----- 4 E2E local wiring -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Deep wiring · local")
    title_bar(s, "Browser → API → SQLite → public edge")
    pic(s, "wiring_e2e.png", Inches(0.25), Inches(1.25), Inches(12.8), Inches(5.6))
    footer(s, 4)

    # ----- 5 Medallion -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Warehouse")
    title_bar(s, "Bronze → Silver → Gold grain")
    pic(s, "arch_flow.png", Inches(0.35), Inches(1.35), Inches(12.6), Inches(3.2))
    details = [
        ("Bronze", "Immutable events. Keys: event_id + (dealer_id, source_id, version)."),
        ("Validate", "Reject bad money/units/channels → quarantine."),
        ("Silver", "Latest valid (dealer_id, sku_id) sell-through + on-hand."),
        ("Gold", "Channel × region × family aggregates."),
        ("Serve", ":8770 /api/metric with session dealer WHERE."),
        ("Edge", "cloudflared tunnel → vc.datasharkbi.com"),
    ]
    for i, (t, b) in enumerate(details):
        col, row = i % 3, i // 3
        left = Inches(0.55) + Inches(col * 4.2)
        top = Inches(4.7) + Inches(row * 1.1)
        add_textbox(s, left, top, Inches(0.95), Inches(0.35), t, size=12, bold=True, color=ACCENT, font="Calibri")
        add_textbox(s, left + Inches(1.0), top, Inches(3.05), Inches(0.95), b, size=11, color=MUTED, font="Calibri")
    footer(s, 5)

    # ----- 6 Snowflake path overview -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Deep wiring · Snowflake")
    title_bar(s, "Optional cloud path — same contracts")
    pic(s, "wiring_snowflake.png", Inches(0.2), Inches(1.2), Inches(12.9), Inches(5.65))
    footer(s, 6)

    # ----- 7 Snowflake ordered setup -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Snowflake setup")
    title_bar(s, "Ordered scripts.cloud pipeline")
    steps = [
        ("doctor", "Verify connection profile (account / user / role / warehouse)"),
        ("platform", "01_platform.sql → DB VC_RETAIL_DEMO + BRONZE/SILVER/GOLD/GOVERNANCE/SERVING + WH"),
        ("load", "Seed local SQLite → PUT JSONL → @BRONZE.INGEST_STAGE → COPY INTO VARIANT"),
        ("transform", "02_transform.sql → normalize → SILVER.FACTS (+ quarantine)"),
        ("governance", "03_governance.sql → USER_DEALERS + row-access policy on SERVING views"),
        ("semantic", "04_semantic.sql → SERVING.VC_RETAIL_SEMANTICS for Cortex Analyst"),
        ("validate", "05_validation.sql smoke checks"),
        ("reconcile", "Compare local silver_facts vs SILVER.FACTS → docs/evidence/reconcile.json"),
    ]
    for i, (cmd, desc) in enumerate(steps):
        y = Inches(1.4) + Inches(i * 0.65)
        add_rect(s, Inches(0.55), y, Inches(2.4), Inches(0.55), fill=DARK)
        add_textbox(s, Inches(0.55), y + Inches(0.1), Inches(2.4), Inches(0.4), cmd, size=13, bold=True, color=LIGHT, font="Calibri", align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(3.15), y + Inches(0.1), Inches(9.5), Inches(0.45), desc, size=13, color=MUTED, font="Calibri")
    footer(s, 7)

    # ----- 8 Snowflake schemas / RAP -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Snowflake security")
    title_bar(s, "Database layout + dealer row policy")
    schemas = [
        ("BRONZE", "DEALERS, PRODUCTS, EVENTS (VARIANT), INGEST_STAGE, INCOMING"),
        ("SILVER", "FACTS (current dealer×SKU), QUARANTINE"),
        ("GOLD", "Channel / family / region rolls"),
        ("GOVERNANCE", "USER_DEALERS map + DEALER_SCOPE row-access policy"),
        ("SERVING", "Secure views for the app reader role + semantic view"),
    ]
    for i, (name, body) in enumerate(schemas):
        y = Inches(1.45) + Inches(i * 0.7)
        add_rect(s, Inches(0.55), y, Inches(2.6), Inches(0.6), fill=ACCENT if i == 3 else DARK)
        add_textbox(s, Inches(0.55), y + Inches(0.12), Inches(2.6), Inches(0.4), name, size=14, bold=True, color=LIGHT, font="Calibri", align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(3.35), y + Inches(0.12), Inches(9.3), Inches(0.45), body, size=14, color=MUTED, font="Calibri")
    add_textbox(
        s,
        Inches(0.55),
        Inches(5.15),
        Inches(12.2),
        Inches(1.5),
        "Row policy: CURRENT_USER() must map in GOVERNANCE.USER_DEALERS to DLR-#### or '*'.\n"
        "Attached on SERVING.SILVER_FACTS, DEALERS, QUARANTINE, GOLD_CHANNEL_FAMILY.\n"
        "Reader role gets USAGE on WH/DB/SERVING + SELECT on serving views — not raw Bronze.",
        size=14,
        color=INK,
        font="Calibri",
    )
    footer(s, 8)

    # ----- 9 Dual backend switch -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "App wiring")
    title_bar(s, "One UI, two Store backends")
    add_rect(s, Inches(0.55), Inches(1.5), Inches(6.0), Inches(5.15), fill=CARD, line=LINE)
    add_rect(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.15), fill=DARK)
    add_textbox(s, Inches(0.85), Inches(1.75), Inches(5.4), Inches(0.4), "Local (default)", size=20, color=INK, font="Georgia")
    add_textbox(s, Inches(7.1), Inches(1.75), Inches(5.4), Inches(0.4), "Snowflake", size=20, color=LIGHT, font="Georgia")
    add_bullets(s, Inches(0.85), Inches(2.4), Inches(5.4), Inches(3.8), [
        "python3 -m app.server",
        "Store.rows → sqlite3 on data/demo.db",
        "Tables: silver_facts, quarantine, …",
        "Dealer WHERE from session cookie",
        "No cloud credentials required",
    ], color=MUTED, size=14)
    add_bullets(s, Inches(7.1), Inches(2.4), Inches(5.4), Inches(3.8), [
        "python3 -m app.server --backend snowflake --connection <reader>",
        "Store.rows → snowflake.connector",
        "USE SCHEMA SERVING",
        "Same metric SQL shapes against views",
        "RAP enforces dealer scope in-warehouse",
    ], color=LIGHT, size=14)
    footer(s, 9)

    # ----- 10 Semantic / Cortex -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Optional Cortex")
    title_bar(s, "Semantic view — not arbitrary SQL")
    add_textbox(
        s,
        Inches(0.65),
        Inches(1.45),
        Inches(12),
        Inches(0.7),
        "SERVING.VC_RETAIL_SEMANTICS binds dimensions (dealer, channel, family, region) and metrics (units, net_sales, margin, on_hand) over SILVER_FACTS.",
        size=15,
        color=MUTED,
        font="Calibri",
    )
    add_bullets(s, Inches(0.65), Inches(2.3), Inches(12), Inches(3.5), [
        "Comment on the view: synthetic VC-flavored snapshot — do not invent forecasts or live order status",
        "scripts/cortex.py calls Cortex Analyst against that semantic view (needs SNOWFLAKE_HOST + PAT)",
        "Offline /api/ask Proposed path remains a closed METRICS vocabulary — not an LLM",
        "Only claim “live Snowflake / Cortex” after scripts.cloud reconcile evidence exists",
        "Isolated DB VC_RETAIL_DEMO — do not reuse other demo databases",
    ], color=MUTED, size=15)
    add_link_button(s, Inches(0.65), Inches(5.9), Inches(5.5), Inches(0.65), "Open public demo →", DEMO_URL, size=16)
    footer(s, 10)

    # ----- 11 Metrics -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Contracts")
    title_bar(s, "Governed metrics (closed vocabulary)")
    pic(s, "wiring_metrics.png", Inches(0.45), Inches(1.3), Inches(12.4), Inches(5.5))
    footer(s, 11)

    # ----- 12 API map -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "HTTP surface")
    title_bar(s, "API map the UI calls")
    apis = [
        ("POST /api/login", "Session cookie; binds dealer or operator"),
        ("GET /api/session · /api/health", "Identity + backend label + synthetic flag"),
        ("GET /api/summary · /api/metric", "Overview KPIs; one of six METRICS"),
        ("GET /api/analytics", "Bundled channel/family/region/stock_risk"),
        ("POST /api/ask", "mode=today|proposed → FAQ / metric / escalate"),
        ("POST /api/escalate", "Structured handoff JSON (demo)"),
        ("GET /api/products · /api/catalog", "Catalog search + metric lineage notes"),
    ]
    for i, (path, desc) in enumerate(apis):
        y = Inches(1.45) + Inches(i * 0.7)
        add_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.6), fill=CARD if i % 2 == 0 else WARM, line=LINE)
        add_textbox(s, Inches(0.7), y + Inches(0.12), Inches(5.0), Inches(0.4), path, size=14, bold=True, color=INK, font="Calibri")
        add_textbox(s, Inches(5.9), y + Inches(0.12), Inches(6.6), Inches(0.4), desc, size=14, color=MUTED, font="Calibri")
    footer(s, 12)

    # ----- 13 Ask routing -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Assistant")
    title_bar(s, "Today vs Proposed request router")
    pic(s, "wiring_ask.png", Inches(0.45), Inches(1.25), Inches(12.4), Inches(5.55))
    footer(s, 13)

    # ----- 14 Gap -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Chatbot gap")
    title_bar(s, "What changes for the dealer")
    pic(s, "gap_split.png", Inches(0.55), Inches(1.3), Inches(12.2), Inches(5.5))
    footer(s, 14)

    # ----- 15 Analytics visuals -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Analytics")
    title_bar(s, "Example panels (synthetic)")
    pic(s, "chart_channel.png", Inches(0.4), Inches(1.35), Inches(6.2), Inches(5.4))
    pic(s, "chart_stock.png", Inches(6.7), Inches(1.35), Inches(6.2), Inches(5.4))
    footer(s, 15)

    # ----- 16 Prompts + URL -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H)
    eyebrow(s, "Demo script")
    title_bar(s, "Prompts that exercise the wiring")
    add_link_button(s, Inches(0.65), Inches(1.4), Inches(12.0), Inches(0.7), f"Open demo  ·  {DEMO_URL.rstrip('/')}", DEMO_URL, fill=DARK, size=16)
    add_rect(s, Inches(0.55), Inches(2.4), Inches(6.0), Inches(4.2), fill=CARD, line=LINE)
    add_rect(s, Inches(6.8), Inches(2.4), Inches(6.0), Inches(4.2), fill=DARK)
    add_textbox(s, Inches(0.85), Inches(2.6), Inches(5.4), Inches(0.4), "Today", size=18, color=INK, font="Georgia")
    add_textbox(s, Inches(7.1), Inches(2.6), Inches(5.4), Inches(0.4), "Proposed", size=18, color=LIGHT, font="Georgia")
    add_bullets(s, Inches(0.85), Inches(3.2), Inches(5.4), Inches(3.0), [
        "custom order lead time → FAQ + source",
        "what is open box → policy",
        "where is my order 12345 → escalate",
    ], color=MUTED, size=14)
    add_bullets(s, Inches(7.1), Inches(3.2), Inches(5.4), Inches(3.0), [
        "show sales by channel → SQL + rows",
        "show stock risk → stock_risk",
        "alabaster → catalog",
        "invent forecast → refuse",
    ], color=LIGHT, size=14)
    footer(s, 16)

    # ----- 17 Close -----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, W, H, DARK)
    add_rect(s, 0, 0, Inches(7.0), H, fill=DARK)
    add_rect(s, Inches(7.0), 0, Inches(0.08), H, fill=ACCENT)
    pic(s, "vc-hero-lighting.jpg", Inches(7.08), 0, Inches(6.25), Inches(7.5))
    add_textbox(s, Inches(0.65), Inches(1.5), Inches(6), Inches(0.35), "LOCAL  ·  SNOWFLAKE  ·  PUBLIC", size=12, bold=True, color=ACCENT, font="Calibri")
    add_textbox(s, Inches(0.65), Inches(2.05), Inches(6), Inches(1.3), "Open the wiring.\nClick the demo.", size=36, color=LIGHT, font="Georgia")
    add_link_button(s, Inches(0.65), Inches(3.7), Inches(5.5), Inches(0.75), "Open live demo →", DEMO_URL, size=18)
    add_link_text(s, Inches(0.65), Inches(4.65), Inches(6), Inches(0.35), DEMO_URL, DEMO_URL, size=14, color=SOFT, bold=False)
    add_textbox(s, Inches(0.65), Inches(5.2), Inches(6), Inches(0.8), "dlr-0001 / vc-demo\nNot Visual Comfort production data.", size=14, color=SOFT, font="Calibri")
    add_textbox(s, Inches(0.65), Inches(6.3), Inches(6), Inches(0.35), "Paul Rendleman", size=15, color=LIGHT, font="Georgia")

    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    primary = "docs/deck/VC_Retail_Analytics_Wiring.pptx"
    print(build(primary))
    print(build("docs/deck/VC_Retail_Analytics_Interview.pptx"))
